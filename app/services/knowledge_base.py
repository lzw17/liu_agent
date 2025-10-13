"""Knowledge base service for PDF processing and vector storage."""
import os
import uuid
import hashlib
from typing import List, Dict, Any, Optional, Tuple
from pathlib import Path
import asyncio
import aiofiles
import PyPDF2
import pdfplumber
from docx import Document
import chromadb
from sentence_transformers import SentenceTransformer
from ..core import get_config, logger
from ..models.schemas import KnowledgeBaseDocument, DocumentUploadResponse
from datetime import datetime


class DocumentProcessor:
    """Document processing utilities."""
    
    @staticmethod
    def extract_text_from_pdf(file_path: str) -> Tuple[str, int]:
        """Extract text from PDF file."""
        text = ""
        pages = 0
        
        try:
            # Try with pdfplumber first (better for complex layouts)
            with pdfplumber.open(file_path) as pdf:
                pages = len(pdf.pages)
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        except Exception as e:
            logger.warning(f"pdfplumber failed, trying PyPDF2: {e}")
            try:
                # Fallback to PyPDF2
                with open(file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    pages = len(pdf_reader.pages)
                    for page in pdf_reader.pages:
                        text += page.extract_text() + "\n"
            except Exception as e2:
                logger.error(f"PDF extraction failed: {e2}")
                raise
        
        return text.strip(), pages
    
    @staticmethod
    def extract_text_from_docx(file_path: str) -> Tuple[str, int]:
        """Extract text from DOCX file."""
        try:
            doc = Document(file_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # Estimate pages (assuming ~500 words per page)
            word_count = len(text.split())
            pages = max(1, word_count // 500)
            
            return text.strip(), pages
        except Exception as e:
            logger.error(f"DOCX extraction failed: {e}")
            raise
    
    @staticmethod
    async def extract_text_from_txt(file_path: str) -> Tuple[str, int]:
        """Extract text from TXT file."""
        try:
            async with aiofiles.open(file_path, 'r', encoding='utf-8') as file:
                text = await file.read()
            
            # Estimate pages
            word_count = len(text.split())
            pages = max(1, word_count // 500)
            
            return text.strip(), pages
        except Exception as e:
            logger.error(f"TXT extraction failed: {e}")
            raise
    
    @staticmethod
    def chunk_text(text: str, chunk_size: int = 1000, chunk_overlap: int = 200) -> List[str]:
        """Split text into overlapping chunks."""
        if len(text) <= chunk_size:
            return [text]
        
        chunks = []
        start = 0
        
        while start < len(text):
            end = start + chunk_size
            
            # Try to break at sentence boundary
            if end < len(text):
                # Look for sentence endings within the last 100 characters
                last_period = text.rfind('.', start + chunk_size - 100, end)
                last_exclamation = text.rfind('!', start + chunk_size - 100, end)
                last_question = text.rfind('?', start + chunk_size - 100, end)
                
                sentence_end = max(last_period, last_exclamation, last_question)
                if sentence_end > start:
                    end = sentence_end + 1
            
            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)
            
            start = end - chunk_overlap
            if start >= len(text):
                break
        
        return chunks


class KnowledgeBaseService:
    """Service for managing document knowledge base."""
    
    def __init__(self):
        self.config = get_config()
        self.processor = DocumentProcessor()
        self.embedding_model = None
        self.chroma_client = None
        self.collection = None
        self._initialize()
    
    def _initialize(self):
        """Initialize the knowledge base service."""
        try:
            # Initialize embedding model
            model_name = self.config.knowledge_base.embedding_model
            logger.info(f"Loading embedding model: {model_name}")
            self.embedding_model = SentenceTransformer(model_name)
            
            # Initialize ChromaDB
            db_path = Path(self.config.knowledge_base.vector_db_path)
            db_path.mkdir(parents=True, exist_ok=True)
            
            self.chroma_client = chromadb.PersistentClient(path=str(db_path))
            self.collection = self.chroma_client.get_or_create_collection(
                name="documents",
                metadata={"hnsw:space": "cosine"}
            )
            
            logger.info("Knowledge base service initialized successfully")
            
        except Exception as e:
            logger.error(f"Failed to initialize knowledge base service: {e}")
            raise
    
    async def upload_document(
        self, 
        file_path: str, 
        filename: str
    ) -> DocumentUploadResponse:
        """Process and store a document in the knowledge base."""
        try:
            # Generate document ID
            doc_id = str(uuid.uuid4())
            
            # Extract text based on file type
            file_ext = Path(filename).suffix.lower()
            
            if file_ext == '.pdf':
                text, pages = self.processor.extract_text_from_pdf(file_path)
            elif file_ext == '.docx':
                text, pages = self.processor.extract_text_from_docx(file_path)
            elif file_ext == '.txt':
                text, pages = await self.processor.extract_text_from_txt(file_path)
            else:
                raise ValueError(f"Unsupported file type: {file_ext}")
            
            if not text.strip():
                raise ValueError("No text content found in document")
            
            # Chunk the text
            chunks = self.processor.chunk_text(
                text,
                self.config.knowledge_base.chunk_size,
                self.config.knowledge_base.chunk_overlap
            )
            
            # Generate embeddings
            logger.info(f"Generating embeddings for {len(chunks)} chunks")
            embeddings = self.embedding_model.encode(chunks).tolist()
            
            # Store in ChromaDB
            chunk_ids = [f"{doc_id}_{i}" for i in range(len(chunks))]
            
            # Get file size
            file_size = os.path.getsize(file_path)
            
            # Current timestamp
            upload_date = datetime.now().isoformat()
            
            metadatas = [
                {
                    "document_id": doc_id,
                    "filename": filename,
                    "chunk_index": i,
                    "total_chunks": len(chunks),
                    "pages": pages,
                    "file_size": file_size,
                    "upload_date": upload_date
                }
                for i in range(len(chunks))
            ]
            
            self.collection.add(
                ids=chunk_ids,
                embeddings=embeddings,
                documents=chunks,
                metadatas=metadatas
            )
            
            # Get file size
            file_size = os.path.getsize(file_path)
            
            logger.info(f"Document {filename} processed successfully: {len(chunks)} chunks")
            
            return DocumentUploadResponse(
                document_id=doc_id,
                filename=filename,
                pages=pages,
                chunks=len(chunks),
                status="success"
            )
            
        except Exception as e:
            logger.error(f"Document processing failed: {e}")
            raise
    
    async def search_documents(
        self, 
        query: str, 
        max_results: Optional[int] = None
    ) -> List[Dict[str, Any]]:
        """Search documents using semantic similarity."""
        try:
            if max_results is None:
                max_results = self.config.knowledge_base.max_results
            
            # Generate query embedding
            query_embedding = self.embedding_model.encode([query]).tolist()[0]
            
            # Search in ChromaDB
            results = self.collection.query(
                query_embeddings=[query_embedding],
                n_results=max_results,
                include=["documents", "metadatas", "distances"]
            )
            
            # Format results
            search_results = []
            if results['documents'] and results['documents'][0]:
                for i, (doc, metadata, distance) in enumerate(zip(
                    results['documents'][0],
                    results['metadatas'][0],
                    results['distances'][0]
                )):
                    search_results.append({
                        "content": doc,
                        "metadata": metadata,
                        "similarity_score": 1 - distance,  # Convert distance to similarity
                        "source": metadata.get("filename", "Unknown")
                    })
            
            return search_results
            
        except Exception as e:
            logger.error(f"Document search failed: {e}")
            return []
    
    async def get_documents(self) -> List[KnowledgeBaseDocument]:
        """Get list of all documents in the knowledge base."""
        try:
            # Get all documents from ChromaDB
            try:
                results = self.collection.get(include=["metadatas"])
                if not results or not results.get('metadatas') or len(results.get('metadatas', [])) == 0:
                    logger.warning("No documents found in knowledge base")
                    return []
            except Exception as e:
                logger.error(f"Failed to get documents from ChromaDB: {e}")
                return []
            
            # Group by document_id
            docs_dict = {}
            for metadata in results['metadatas']:
                doc_id = metadata['document_id']
                if doc_id not in docs_dict:
                    # Parse upload date from ISO format string
                    upload_date = None
                    if 'upload_date' in metadata:
                        try:
                            upload_date = datetime.fromisoformat(metadata['upload_date'])
                        except (ValueError, TypeError):
                            upload_date = datetime.now()
                    else:
                        upload_date = datetime.now()
                        
                    docs_dict[doc_id] = {
                        'id': doc_id,
                        'filename': metadata['filename'],
                        'chunks': 0,
                        'pages': metadata.get('pages', 0),
                        'file_size': metadata.get('file_size', 0),
                        'upload_date': upload_date
                    }
                docs_dict[doc_id]['chunks'] += 1
            
            # Convert to KnowledgeBaseDocument objects
            documents = []
            for doc_data in docs_dict.values():
                documents.append(KnowledgeBaseDocument(
                    id=doc_data['id'],
                    filename=doc_data['filename'],
                    upload_date=doc_data['upload_date'],
                    pages=doc_data['pages'],
                    chunks=doc_data['chunks'],
                    file_size=doc_data['file_size']
                ))
            
            return documents
            
        except Exception as e:
            logger.error(f"Failed to get documents: {e}")
            return []
    
    async def delete_document(self, document_id: str) -> bool:
        """Delete a document from the knowledge base."""
        try:
            # Get all chunk IDs for this document
            results = self.collection.get(
                where={"document_id": document_id},
                include=["metadatas"]
            )
            
            if not results['ids']:
                return False
            
            # Delete all chunks
            self.collection.delete(ids=results['ids'])
            
            logger.info(f"Document {document_id} deleted successfully")
            return True
            
        except Exception as e:
            logger.error(f"Failed to delete document {document_id}: {e}")
            return False


# Global knowledge base service instance
_kb_service: Optional[KnowledgeBaseService] = None


def get_knowledge_base_service() -> KnowledgeBaseService:
    """Get the global knowledge base service instance."""
    global _kb_service
    if _kb_service is None:
        _kb_service = KnowledgeBaseService()
    return _kb_service
